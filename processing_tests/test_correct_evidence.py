"""Check correction mechanics with disposable fixtures and read-only inputs."""
import contextlib
import copy
import io
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from openpyxl.styles import Font

SOURCE = Path(__file__).parents[1] / 'tools' / 'correct-evidence.py'
spec = importlib.util.spec_from_file_location('correction_under_test', SOURCE)
correction = importlib.util.module_from_spec(spec)
spec.loader.exec_module(correction)


class CorrectionMechanicsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.site = Path(self.temp.name)
        (self.site / 'sources').mkdir()
        (self.site / 'sources' / 'fixture.txt').write_bytes(b'Original archive\n')
        (self.site / 'families').mkdir()
        self.site_patch = patch.object(correction, 'SITE', self.site)
        self.site_patch.start()
        self.addCleanup(self.site_patch.stop)
        self.finding = dict(field='trunk_km', original='=C11/E11 *1000', status='supported',
                            supported_value=0, evidence='Synthetic source boundary.',
                            selected_observation='O1', source_ids='S1', pages='1')
        self.project = dict(family_slug='fixture', project_id='P1',
                            summary=dict(name='Fixture project', analysis_conclusion='Qualified conclusion'),
                            policy_chart=dict(original_row=dict(project='Fixture chart', trunk_km=101),
                                              overall_status='Qualified status', findings=[self.finding]),
                            observations=[dict(id='O1', project_id='P1', measure='length', value=0,
                                               unit='km', excerpt='Raw | excerpt\nretained',
                                               source_id='S1', pages='1', notes='Original notes')],
                            components=[dict(id='C1', stage='Fixture stage', category='Synthetic', description='Fixture component', boundary='Component boundary', source_id='S1', pages='1')],
                            timeline=[dict(date='2026', event='Fixture event', status='Fixture status', detail='Fixture detail', source_id='S1', pages='1')],
                            warnings=[dict(severity='info', code='FIXTURE', finding='Qualified warning', implication='Fixture qualification', source_ids='S1', pinpoints='1')])
        self.pack = dict(projects=[self.project], sources=[dict(source_id='S1', manifest_included=True)],
                         generated=dict(research_cutoff='2026-09-04'), agent_instructions=[])
        self.initial_pack = copy.deepcopy(self.pack)
        self.wb = openpyxl.Workbook()
        self.wb.active.title = 'Policy Paper Chart'
        for name in ['Chart Review', 'Projects', 'Observations', 'Components', 'Timeline', 'Warnings', 'Sources']:
            self.wb.create_sheet(name)
        self.wb['Policy Paper Chart']['A5'] = 'Fixture chart'
        self.wb['Policy Paper Chart']['N5'] = '=C11/E11 *1000'
        self.wb['Policy Paper Chart']['N5'].font = Font(bold=True, color='FF123456')
        self.wb['Policy Paper Chart']['O5'] = 'Old status'
        self.wb['Chart Review']['A5'] = 'Fixture project'
        self.wb['Chart Review']['C5'] = 'trunk_km'
        self.wb['Chart Review']['D5'] = '=C11/E11 *1000'
        self.wb['Projects']['A5'] = 'fixture'
        self.wb['Observations']['C5'] = 'O1'
        self.wb['Components']['B5'] = 'C1'
        for sheet, columns in [('Chart Review', 10), ('Projects', 15), ('Observations', 19), ('Components', 8)]:
            self.wb[sheet].cell(5, columns)

    def test_correction_qualifies_zero_without_changing_raw_evidence_or_cutoff(self):
        original = copy.deepcopy(self.project['policy_chart']['original_row'])
        raw = self.project['observations'][0]['excerpt']
        correction.correct(self.pack)
        self.assertEqual(self.finding['supported_value'], '0 km. Boundary/qualification: Synthetic source boundary.')
        self.assertEqual(self.finding['original'], '=C11/E11 *1000')
        self.assertEqual(self.project['policy_chart']['original_row'], original)
        self.assertEqual(self.project['observations'][0]['excerpt'], raw)
        self.assertEqual(self.pack['generated']['research_cutoff'], '2026-09-04')
        self.assertEqual(self.pack['generated']['audit_correction_date'], '2026-09-07')
        self.assertFalse(self.pack['sources'][0]['manifest_included'])
        self.assertEqual((self.site / 'sources' / 'fixture.txt').read_bytes(), b'Original archive\n')

    def test_correction_is_idempotent_and_does_not_duplicate_guidance(self):
        correction.correct(self.pack)
        before = copy.deepcopy(self.pack)
        correction.correct(self.pack)
        self.assertEqual(self.pack, before)
        self.assertEqual(len(self.pack['agent_instructions']), 1)

    def test_actual_corrected_snapshot_is_unchanged_in_memory(self):
        path = SOURCE.parents[1] / 'site' / 'evidence-pack.json'
        original = path.read_bytes()
        pack = json.loads(original)
        expected = json.dumps(pack, sort_keys=True, ensure_ascii=False, allow_nan=False)
        for _ in range(2):
            correction.correct(pack)
            self.assertEqual(json.dumps(pack, sort_keys=True, ensure_ascii=False, allow_nan=False), expected)
        self.assertEqual(path.read_bytes(), original)

    def test_protected_guard_detects_each_immutable_input_change(self):
        baseline = correction.protected(self.pack, self.wb)
        self.assertEqual(baseline['sources'], {'sources/fixture.txt': hashlib.sha256(b'Original archive\n').hexdigest()})
        self.assertEqual(set(baseline), {'sources', 'original_rows', 'original_findings', 'excerpts', 'chart_inputs'})
        for key, change in [
            ('original_rows', lambda: self.project['policy_chart']['original_row'].update(trunk_km=102)),
            ('original_findings', lambda: self.finding.update(original='Changed input')),
            ('excerpts', lambda: self.project['observations'][0].update(excerpt='Changed excerpt')),
            ('chart_inputs', lambda: setattr(self.wb['Policy Paper Chart']['N5'], 'value', '=1+1')),
            ('chart_inputs', lambda: setattr(self.wb['Policy Paper Chart']['N5'], 'font', Font(italic=True))),
        ]:
            with self.subTest(key=key):
                pack_before = copy.deepcopy(self.pack)
                cell = self.wb['Policy Paper Chart']['N5']
                value, font = cell.value, copy.copy(cell.font)
                change()
                self.assertNotEqual(correction.protected(self.pack, self.wb)[key], baseline[key])
                self.pack.clear(); self.pack.update(pack_before)
                self.project = self.pack['projects'][0]
                self.finding = self.project['policy_chart']['findings'][0]
                cell.value, cell.font = value, font
        self.wb['Policy Paper Chart']['O5'] = 'Permitted audit annotation'
        self.assertEqual(correction.protected(self.pack, self.wb), baseline)

    def test_source_guard_tracks_bytes_additions_and_removals(self):
        path = self.site / 'sources' / 'fixture.txt'
        expected = {'sources/fixture.txt': hashlib.sha256(b'Original archive\n').hexdigest()}
        self.assertEqual(correction.sources(), expected)
        path.write_bytes(b'Changed archive\n')
        self.assertNotEqual(correction.sources(), expected)
        extra = self.site / 'sources' / 'extra.txt'
        extra.write_bytes(b'Extra archive')
        self.assertEqual(set(correction.sources()), {'sources/fixture.txt', 'sources/extra.txt'})
        path.unlink()
        self.assertEqual(correction.sources(), {'sources/extra.txt': hashlib.sha256(b'Extra archive').hexdigest()})

    def test_workbook_updates_replicas_and_keeps_original_formula_literal_in_audit(self):
        self.wb['Sources']['B5'] = 'A long synthetic research sentence to replace in replicas.'
        self.wb['Policy Paper Chart']['B5'] = self.wb['Sources']['B5'].value
        self.wb['Chart Review']['K5'] = self.wb['Sources']['B5'].value
        self.wb['Chart Review']['A6'] = 'Unmatched original audit row'
        self.wb['Chart Review']['D6'] = self.wb['Sources']['B5'].value
        before = {(cell.coordinate, cell.value, cell.data_type, str(cell._style))
                  for row in self.wb['Policy Paper Chart'] for cell in row if cell.column <= 14}
        correction.correct(self.pack)
        correction.update_workbook(self.wb, {}, self.pack, {
            'A long synthetic research sentence': 'Shorter overlapping replacement',
            'A long synthetic research sentence to replace in replicas.': 'Qualified replica text',
        })
        after = {(cell.coordinate, cell.value, cell.data_type, str(cell._style))
                 for row in self.wb['Policy Paper Chart'] for cell in row if cell.column <= 14}
        self.assertEqual(after, before)
        self.assertEqual(self.wb['Chart Review']['D5'].value, '=C11/E11 *1000')
        self.assertEqual(self.wb['Chart Review']['D5'].data_type, 's')
        self.assertEqual(self.wb['Chart Review']['F5'].value, self.finding['supported_value'])
        self.assertEqual(self.wb['Chart Review']['K5'].value, 'Qualified replica text')
        self.assertEqual(self.wb['Chart Review']['D6'].value, 'A long synthetic research sentence to replace in replicas.')
        self.assertEqual(self.wb['Sources']['B5'].value, 'Qualified replica text')
        self.assertEqual(self.wb['Observations']['G5'].value, 0)
        self.assertEqual(self.wb['Observations']['Q5'].value, 'Raw | excerpt\nretained')
        self.assertEqual(self.wb['Observations']['R5'].value, 'selected')
        self.assertEqual(self.wb['Components']['F5'].value, 'Component boundary')
        self.assertEqual(self.wb['Timeline']['C5'].value, 'Fixture event')
        self.assertEqual(self.wb['Warnings']['C5'].value, 'FIXTURE')
        self.assertEqual(self.wb['Policy Paper Chart']['O5'].value, 'Qualified status')
        self.assertEqual(self.wb['Projects']['N5'].value, 'Qualified conclusion')
        self.finding['selected_observation'] = None
        correction.update_workbook(self.wb, {}, self.pack, {})
        self.assertIsNone(self.wb['Observations']['R5'].value)

    def valid_synthetic_export(self):
        self.pack = copy.deepcopy(self.initial_pack)
        self.project = self.pack['projects'][0]
        self.finding = self.project['policy_chart']['findings'][0]
        fields = ['cost_aud_billion', 'network_capacity_gw', 'trunk_km',
                  'millions_per_trunk_km', 'circuits', 'lines', 'syncons_mva',
                  'hubs', 'transformers', 'feeder_lines_km', 'date', 'alt_routes']
        template = copy.deepcopy(self.project)
        self.pack['projects'] = []
        self.wb = openpyxl.Workbook()
        self.wb.active.title = 'Policy Paper Chart'
        self.wb.active.cell(5, 15)
        for name, columns in [('Chart Review', 10), ('Projects', 15), ('Observations', 19),
                              ('Components', 8), ('Timeline', 7), ('Warnings', 7), ('Sources', 6)]:
            self.wb.create_sheet(name).cell(5, columns)
        for number in range(10):
            project = copy.deepcopy(template)
            project['family_slug'] = f'fixture-{number}'
            project['summary']['name'] = f'Fixture project {number}'
            project['policy_chart']['original_row']['project'] = f'Fixture chart {number}'
            project['observations'][0]['id'] = f'O{number}'
            project['components'][0]['id'] = f'C{number}'
            project['policy_chart']['findings'] = []
            for offset, field in enumerate(fields):
                finding = copy.deepcopy(self.finding)
                finding.update(field=field, selected_observation=f'O{number}', supported_value=None)
                project['policy_chart']['findings'].append(finding)
                row = 5 + 12 * number + offset
                self.wb['Chart Review'].cell(row, 1, project['summary']['name'])
                self.wb['Chart Review'].cell(row, 3, field)
            self.pack['projects'].append(project)
            for sheet, column, value in [('Projects', 1, project['family_slug']),
                                          ('Observations', 3, f'O{number}'),
                                          ('Components', 2, f'C{number}'),
                                          ('Policy Paper Chart', 1, f'Fixture chart {number}')]:
                self.wb[sheet].cell(5 + number, column, value)
            self.wb['Policy Paper Chart'].cell(5 + number, 14, '=1+1')
        correction.correct(self.pack)
        correction.update_workbook(self.wb, {}, self.pack, {})
        self.write_families()

    def write_families(self):
        for project in self.pack['projects']:
            path = self.site / 'families' / (project['family_slug'] + '.json')
            path.write_text(json.dumps({**project, 'sources': self.pack['sources'], 'build': {'research_cutoff': self.pack['generated']['research_cutoff'], 'family_slug': project['family_slug'], 'project_id': project['project_id']}}), encoding='utf-8')
            path.with_suffix('.md').write_text(correction.markdown(project), encoding='utf-8')

    def test_complete_synthetic_validation_accepts_coherent_exports(self):
        self.valid_synthetic_export()
        with contextlib.redirect_stdout(io.StringIO()) as output:
            correction.validate(self.pack, self.wb)
        self.assertIn('PASS: 120', output.getvalue())

    def test_rebuilding_synthetic_export_restores_valid_metadata(self):
        self.valid_synthetic_export()
        self.pack['generated'].update(research_cutoff='2026-10-22', path_base='../')
        self.valid_synthetic_export()
        with contextlib.redirect_stdout(io.StringIO()):
            correction.validate(self.pack, self.wb)

    def test_complete_validation_rejects_drift_in_each_replica_and_contract(self):
        for label, damage in [
            ('research cutoff', lambda: self.pack['generated'].update(research_cutoff='2026-10-22')),
            ('path base', lambda: self.pack['generated'].update(path_base='../')),
            ('finding workbook', lambda: setattr(self.wb['Chart Review']['G5'], 'value', 'Changed evidence')),
            ('project workbook', lambda: setattr(self.wb['Projects']['N5'], 'value', 'Changed conclusion')),
            ('chart annotation', lambda: setattr(self.wb['Policy Paper Chart']['O5'], 'value', 'Changed status')),
            ('observation workbook', lambda: setattr(self.wb['Observations']['S5'], 'value', 'Changed notes')),
            ('timeline workbook', lambda: setattr(self.wb['Timeline']['C5'], 'value', 'Changed event')),
            ('warning workbook', lambda: setattr(self.wb['Warnings']['C5'], 'value', 'Changed code')),
            ('selection marker', lambda: setattr(self.wb['Observations']['R5'], 'value', None)),
            ('literal original formula', lambda: setattr(self.wb['Chart Review']['D45'], 'data_type', 'f')),
        ]:
            with self.subTest(label=label):
                self.valid_synthetic_export()
                with contextlib.redirect_stdout(io.StringIO()):
                    correction.validate(self.pack, self.wb)
                damage()
                with self.assertRaises(AssertionError), contextlib.redirect_stdout(io.StringIO()):
                    correction.validate(self.pack, self.wb)
        self.valid_synthetic_export()
        with contextlib.redirect_stdout(io.StringIO()):
            correction.validate(self.pack, self.wb)
        path = self.site / 'families' / 'fixture-0.json'
        data = json.loads(path.read_text())
        data['observations'][0]['value'] = 123
        path.write_text(json.dumps(data), encoding='utf-8')
        with self.assertRaises(AssertionError):
            correction.validate(self.pack, self.wb)
        self.valid_synthetic_export()
        with contextlib.redirect_stdout(io.StringIO()):
            correction.validate(self.pack, self.wb)
        path.with_suffix('.md').write_text('Unreconciled Markdown', encoding='utf-8')
        with self.assertRaises(AssertionError):
            correction.validate(self.pack, self.wb)

    def test_check_entry_point_is_read_only_and_rejects_changed_archived_bytes(self):
        self.valid_synthetic_export()
        (self.site / 'evidence-pack.json').write_text(json.dumps(self.pack), encoding='utf-8')
        self.wb.save(self.site / 'evidence-workbook.xlsx')
        guard = correction.protected(self.pack, openpyxl.load_workbook(self.site / 'evidence-workbook.xlsx'))
        log = self.site / 'fixture-log.json'
        log.write_text(json.dumps(dict(protected_before=guard, protected_after=guard)), encoding='utf-8')
        before = {path.relative_to(self.site): path.read_bytes() for path in self.site.rglob('*') if path.is_file()}
        with patch.object(correction, 'LOG', log), patch.object(sys, 'argv', ['correct-evidence.py', '--check']), contextlib.redirect_stdout(io.StringIO()):
            correction.main()
        self.assertEqual({path.relative_to(self.site): path.read_bytes() for path in self.site.rglob('*') if path.is_file()}, before)
        (self.site / 'sources' / 'fixture.txt').write_bytes(b'Changed archived bytes')
        with patch.object(correction, 'LOG', log), patch.object(sys, 'argv', ['correct-evidence.py', '--check']), contextlib.redirect_stdout(io.StringIO()), self.assertRaises(AssertionError):
            correction.main()

    def test_write_entry_point_changes_only_disposable_exports_and_records_its_delta(self):
        self.valid_synthetic_export()
        self.pack['generated']['path_base'] = 'historical-fixture/'
        (self.site / 'evidence-pack.json').write_text(json.dumps(self.pack), encoding='utf-8')
        self.wb.save(self.site / 'evidence-workbook.xlsx')
        workbook = openpyxl.load_workbook(self.site / 'evidence-workbook.xlsx')
        before = correction.protected(self.pack, workbook)
        source_bytes = (self.site / 'sources' / 'fixture.txt').read_bytes()
        log = self.site / 'new-fixture-log.json'
        with patch.object(correction, 'LOG', log), patch.object(sys, 'argv', ['correct-evidence.py']), contextlib.redirect_stdout(io.StringIO()):
            correction.main()
        actual = json.loads((self.site / 'evidence-pack.json').read_text())
        actual_workbook = openpyxl.load_workbook(self.site / 'evidence-workbook.xlsx')
        self.assertEqual(actual['generated']['path_base'], './')
        self.assertEqual(actual['generated']['research_cutoff'], '2026-09-04')
        self.assertEqual(correction.protected(actual, actual_workbook), before)
        self.assertEqual((self.site / 'sources' / 'fixture.txt').read_bytes(), source_bytes)
        delta = json.loads(log.read_text())
        self.assertEqual(delta['protected_before'], before)
        self.assertEqual(delta['protected_after'], before)
        self.assertIn(dict(path='/generated/path_base', before='historical-fixture/', after='./'), delta['changes'])
        self.assertEqual(delta['family_by_json_path']['/projects/0'], 'fixture-0')
        self.assertEqual(delta['record_ids_by_json_path']['/projects/0/observations/0'], 'O0')
        with contextlib.redirect_stdout(io.StringIO()):
            correction.validate(actual, actual_workbook)

    def test_change_log_retains_nested_paths_and_before_after_values(self):
        changes = list(correction.changes({'records': [dict(value=0)]}, {'records': [dict(value=2)]}))
        self.assertEqual(changes, [dict(path='/records/0/value', before=0, after=2)])

    def test_markdown_preserves_raw_text_and_escapes_table_delimiters(self):
        family = self.site / 'families' / 'fixture.json'
        family.write_text(json.dumps(dict(sources=[dict(source_id='S1', title='Source | title')])) , encoding='utf-8')
        before = family.read_bytes()
        output = correction.markdown(self.project)
        self.assertIn('Raw \\| excerpt retained', output)
        self.assertIn('Source \\| title', output)
        self.assertIn('Research cutoff: 2026-09-04. Archived-evidence audit correction: 2026-09-07 (not new research).', output)
        self.assertIn('"trunk_km": 101', output)
        self.assertEqual(self.project['observations'][0]['excerpt'], 'Raw | excerpt\nretained')
        self.assertEqual(family.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
