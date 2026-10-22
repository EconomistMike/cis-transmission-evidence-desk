"""Exercise the existing presentation synchroniser only in disposable fixture trees."""
import copy
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

import openpyxl

SOURCE = Path(__file__).parents[1] / 'tools' / 'embed-evidence.py'


class EmbedEvidenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.site = self.root / 'site'
        self.site.mkdir()
        (self.root / 'tools').mkdir()
        self.script = self.root / 'tools' / 'embed-evidence.py'
        shutil.copyfile(SOURCE, self.script)
        self.original_row = dict(project='Fixture project', cost=12.5)
        self.project = dict(family_slug='fixture', summary=dict(name='Fixture <scope> & “name”'),
                            policy_chart=dict(original_row=copy.deepcopy(self.original_row)),
                            observations=[dict(id='O1', notes='</script><script>synthetic</script>')])
        self.pack = dict(projects=[self.project], sources=[dict(source_id='S1')],
                         build_report=dict(status='fixture'),
                         generated=dict(research_cutoff='2026-09-04', audit_correction_date='2026-09-07'))
        self.old = dict(packs=[dict(family_slug='fixture', policy_chart=dict(original_row=copy.deepcopy(self.original_row)))],
                        sources=[], report={}, build=dict(contract='v1', research_cutoff='2020-01-01', retained='keep'),
                        evidence_columns=['column' + str(n) for n in range(11)])
        self.metadata = [None, 0, 'Class 5', 'Uncertain', 'Proposed', 'Owner', 'Boundary', 500, 'Earlier', 'Dependent', '=A1']
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Policy Paper Chart'
        sheet.cell(5, 1, 'Fixture project')
        for column, value in enumerate(self.metadata, 15):
            sheet.cell(5, column, value)
        workbook.save(self.site / 'evidence-workbook.xlsx')
        self.write_inputs()

    def write_inputs(self):
        (self.site / 'evidence-pack.json').write_text(json.dumps(self.pack), encoding='utf-8')
        self.prefix = '<!doctype html><title>Fixture</title><script id="evidence" type="application/json">'
        self.suffix = '</script><p>Retained shell</p>'
        (self.site / 'evidence-desk.html').write_text(self.prefix + json.dumps(self.old) + self.suffix, encoding='utf-8')

    def run_script(self):
        return subprocess.run([sys.executable, str(self.script)], cwd=self.root, capture_output=True, text=True,
                              timeout=15, env=dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONOPTIMIZE='0'))

    def payload(self):
        text = (self.site / 'evidence-desk.html').read_text(encoding='utf-8')
        match = re.search(r'<script[^>]*id="evidence"[^>]*>(.*?)</script>', text, re.S)
        return json.loads(match[1])

    def test_synchronises_records_and_metadata_without_changing_original_inputs(self):
        inputs = {name: (self.site / name).read_bytes() for name in ['evidence-pack.json', 'evidence-workbook.xlsx']}
        result = self.run_script()
        self.assertEqual(result.returncode, 0, result.stderr)
        actual = self.payload()
        self.assertEqual(actual['packs'], [{**self.project, 'name': self.project['summary']['name']}])
        self.assertEqual(actual['sources'], self.pack['sources'])
        self.assertEqual(actual['report'], self.pack['build_report'])
        self.assertEqual(actual['build'], dict(contract='v1', retained='keep', **self.pack['generated']))
        self.assertEqual(actual['evidence_values'], dict(fixture=['' if value is None else value for value in self.metadata]))
        self.assertEqual(actual['link_base'], '')
        for name, before in inputs.items():
            self.assertEqual((self.site / name).read_bytes(), before, name)

    def test_payload_preserves_text_without_closing_its_script_or_changing_the_shell(self):
        self.assertEqual(self.run_script().returncode, 0)
        text = (self.site / 'evidence-desk.html').read_text(encoding='utf-8')
        self.assertTrue(text.startswith(self.prefix))
        self.assertTrue(text.endswith(self.suffix))
        self.assertEqual(text.count('</script>'), 1)
        self.assertEqual(self.payload()['packs'][0]['observations'], self.project['observations'])
        self.assertEqual(self.payload()['packs'][0]['summary']['name'], self.project['summary']['name'])

    def test_repeated_synchronisation_is_byte_identical(self):
        self.assertEqual(self.run_script().returncode, 0)
        before = (self.site / 'evidence-desk.html').read_bytes()
        self.assertEqual(self.run_script().returncode, 0)
        self.assertEqual((self.site / 'evidence-desk.html').read_bytes(), before)

    def test_changed_original_chart_is_rejected_before_writing_html(self):
        self.project['policy_chart']['original_row']['cost'] = 13.5
        self.write_inputs()
        before = (self.site / 'evidence-desk.html').read_bytes()
        result = self.run_script()
        self.assertNotEqual(result.returncode, 0)
        self.assertTrue(result.stderr.rstrip().endswith('ValueError: Original colleague chart was changed: fixture'), result.stderr)
        self.assertEqual((self.site / 'evidence-desk.html').read_bytes(), before)

    def test_missing_evidence_element_is_rejected_before_writing(self):
        before = b'<html><title>No evidence element</title></html>'
        (self.site / 'evidence-desk.html').write_bytes(before)
        self.assertNotEqual(self.run_script().returncode, 0)
        self.assertEqual((self.site / 'evidence-desk.html').read_bytes(), before)

    def test_misaligned_workbook_metadata_is_rejected_before_writing(self):
        self.old['evidence_columns'].pop()
        self.write_inputs()
        before = (self.site / 'evidence-desk.html').read_bytes()
        self.assertNotEqual(self.run_script().returncode, 0)
        self.assertEqual((self.site / 'evidence-desk.html').read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
