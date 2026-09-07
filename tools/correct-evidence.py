#!/usr/bin/env python3
"""Repeatable archived-evidence corrections. Run from repository root; --check validates only.
No research acquisition, source writes, chart-input edits, git mutations or UI regeneration.
"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / 'site'
LOG = ROOT / 'tools/evidence-correction-log.json'

def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()

def sources():
    return {str(p.relative_to(SITE)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted((SITE / 'sources').rglob('*')) if p.is_file()}

def protected(pack, wb):
    return {'sources': sources(), 'original_rows': digest([p['policy_chart']['original_row'] for p in pack['projects']]),
            'original_findings': digest([[f['original'] for f in p['policy_chart']['findings']] for p in pack['projects']]),
            'excerpts': digest([[o.get('excerpt') for o in p['observations']] for p in pack['projects']]),
            'chart_inputs': digest([(c.coordinate, c.value, c.data_type, str(c._style)) for row in wb['Policy Paper Chart'] for c in row if c.column <= 14])}

MARINUS = ('Project Marinus includes Marinus Link and NWTD (2025 RIT-T pp6,18). The $5,035m Stage 1 and $2,535m Stage 2 figures, in June 2023 prices, are Project Marinus study-period costs, not Link-only construction outturn; the valuation period is shorter than project life and depends on cable 2 timing. AEMO 2026 p59 reports $7,570m in the feedback-loop context. Stage 2 Option 2A differs from Option 2: the source removes $184m Burnie-Hampshire Hills scope transferred to T2 Option 1, giving $2,351m in $2023 and $2,485m in $2025. No like-for-like Link-only replacement for the chart\'s $5.5bn is established; do not derive one by subtraction. Funding is not capex.')
SRS_COST = ('The immutable chart contains $3.5bn, consistent with rounding the PADR Option 6 total of $3,519m nominal to one decimal billion. The source gives Stage 1 $1,135m and Stage 2 $2,384m, P50 Class 5 (+100%/-50%). AEMO\'s $2,621m modelled $2025 amount is a different price basis. No transcription error is established (PADR p38; AEMO 2026 p43).')
SRS_RATIO = ('30.70175438596491 = 3,500 AUD million / 114 km, using the original chart inputs. The $3.5bn rounds the $3,519m nominal Option 6 estimate; the 114 km is the older 2023 preparatory route, whereas the 2026 PADR describes approximately 100 km on a revised scope. This mixed-vintage ratio is not a reconciled source-defined construction cost per km (2023 preparatory p8; PADR pp38,43).')
SPLIT = ('The identical $3.7bn NSW and Victoria chart entries have no evidenced current state allocation. $3.7bn is not half of the 2026 $7.6bn whole-VNI total (half is $3.8bn). The author\'s exact derivation is unresolved. Historical PACR Option 5A allocations ($1,755m Victoria / $1,744m NSW, FY2020-21, p48) do not allocate the 2026 total; NSW $948.8m approved early works is not a NSW whole-project estimate.')
CWO_LENGTH = ('The chart\'s 150 km has positive support: CWO-18 p11 states a corridor of around 150 km and CWO-26 derivative line18 states 150 km of 330 kV generator-connection lines. CWO-23 p10 instead describes approximately 166 km of H2P lines, with conditional separable portions (pp12,19). Different vintages, corridor versus line measurement and conditional scope remain unreconciled. Retain an unresolved source discrepancy, not a same-boundary numerical disproof; five portions were excluded in 2025.')
CWO_TIME = ('2029 is not the AEMO 2026 committed-project timing of December 2028 (CWO-22 p23), but falls within the AER infrastructure-completion window April 2028-April 2030 (CWO-18 p12). Expansion timing is March 2033 (CWO-22 p57). The chart milestone is unspecified; a timing conflict applies only if comparing the same committed-project milestone.')
PEC_RATIO = ('25.625 = 4,100 / 160 uses the chart inputs, but $4.1bn is social value, not capex (PEC-TG-PROJECT-PAGE line1705), and 160 km is the sourced Stage 1 SA-border-Buronga-and-Victoria length (PEC-TG-2026-COMPLETION lines161-166), not the whole 900 km project. A social-benefit numerator and a stage-only denominator do not establish project construction cost per km.')
PEC_LINES = ('AEMO 2024 pp17-18 lists multiple physical sections/lines. The chart\'s one line may instead mean one interconnector corridor; no chart-compatible line count is established without defining lines versus sections/corridors.')
PEC_FEEDER = ('Approximately 22 km of 220 kV Buronga-Red Cliffs works is sourced (AEMO 2024 p18; Transgrid project page line1262). It is an interconnector spur: zero lower-voltage works is false, but whether this is a hub-to-project feeder depends on the chart definition. No unqualified REZ-feeder replacement is established.')


def correct(pack):
    replacements = {}
    def setv(obj, key, value):
        old = obj.get(key)
        if old != value:
            if isinstance(old, str) and old: replacements[old] = value
            obj[key] = value
    def narrative(v, slug):
        if isinstance(v, list):
            for a in v: narrative(a, slug)
        elif isinstance(v, dict):
            for k, s in list(v.items()):
                if k in ('original_row', 'original', 'excerpt'): continue
                if not isinstance(s, str): narrative(s, slug); continue
                new = s
                if slug == 'sydney-ring-south' and re.search(r'3[,.]510|transpos|correct the transcription|unverified transcription', s):
                    new = SRS_RATIO if ('ratio' in s or 'divided' in s or 'per-trunk' in s) else SRS_COST
                if slug.startswith('vni-west') and ('half' in s or 'halved' in s or 'author split' in s):
                    new = SPLIT
                if slug == 'central-west-orana':
                    if 'every archived timing' in s: new = s.replace('the date cell is contradicted by every archived timing', 'the date cell requires a defined milestone and the feeder length has an unresolved source discrepancy')
                    if ('150 km' in s and ('contradict' in s or 'same ' in s)) or 'chart\'s feeder_lines_km finding is recorded as contradicted' in s: new = CWO_LENGTH
                if slug == 'marinus-link':
                    if 'excluding separately stated NWTD' in s: new = s.replace('excluding separately stated NWTD', 'including NWTD; study-period valuation, not lifetime construction outturn')
                    if ('chart\'s own' in s and ('7,570' in s or '7570' in s)) or 'Project Marinus plus NWTD' in s: new = MARINUS
                    if 'NWTD works share scope' in s: new = 'NWTD is a component of Project Marinus. Tasmania REZ Expansion is a separate actionable project receiving the Burnie-Hampshire Hills scope under Option 2A.'
                    if 'Do not allocate Tasmania REZ Expansion or NWTD' in s: new = 'Do not double count NWTD within Project Marinus or re-add transferred T2 Option 1 scope to Option 2A. A Link-only cost is unresolved.'
                    if 'contradicted cost cell' in s: new = s.replace('contradicted cost cell', 'unresolved Link-only cost cell')
                    if 'Until the author decides, the cost cell stays at the source-defined boundary' in s: new = 'The Link-only cost replacement remains unresolved; the trunk cell stays derived and alt_routes and hubs remain author judgments.'
                    if 'cost and capacity cells are contradicted' in s: new = 'The cost replacement is unresolved on a Link-only boundary; two-link nameplate capacity is 1.5 GW, contradicting 1.4 GW. Length and ratio remain author derivations; topology gaps remain.'
                if slug == 'project-energyconnect':
                    if '160 km is one 165 km section' in s: new = PEC_RATIO
                    new = new.replace('capacity, syncon, transformer and feeder cells are contradicted by the archived scope', 'capacity, syncon and transformer cells conflict on stated source boundaries; line and feeder counts require explicit definitions')
                    new = new.replace('Cost, capacity, trunk km, syncons, transformers and feeder km all fail on the chart\'s boundaries.', 'Cost, capacity, trunk km, syncons and transformers require the stated boundary corrections; line and feeder interpretations depend on explicit chart definitions.')
                setv(v, k, new)
    for p in pack['projects']:
        slug = p['family_slug']; narrative(p, slug)
        fs = {f['field']: f for f in p['policy_chart']['findings']}
        obs = {o['id']: o for o in p['observations']}
        def finding(field, evidence, status, value):
            for k, v in [('evidence', evidence), ('status', status), ('supported_value', value)]: setv(fs[field], k, v)
        if slug == 'marinus-link':
            finding('cost_aud_billion', MARINUS, 'gap', None)
            setv(p['summary'], 'analysis_conclusion', MARINUS)
            for w in p['warnings']:
                if w['code']=='CONTRADICTION':setv(w,'finding',MARINUS+' Separately, the chart\'s 1.4 GW contradicts the source-defined two-link nameplate of 1.5 GW (AEMO 2026 pp26,58-59).')
            setv(p['summary'], 'identity', 'Project Marinus comprises Marinus Link (HVDC interconnector) and TasNetworks NWTD (HVAC network works); Stage 1 and Stage 2 and Option 2/2A are distinct scope/version boundaries.')
            for oid in ['ML-AO-2025-STAGE1','ML-AO-2025-STAGE2','ML-AO-2026-FL-TOTAL']:
                setv(obs[oid], 'comparability', MARINUS)
            setv(obs['ML-AO-2026-FL-TOTAL'], 'boundary', 'Project Marinus Stage 1 plus Stage 2 including NWTD, feedback-loop reported study-period cost; not Link-only')
            for o in p['observations']:
                if o['id'].startswith('ML-AO-2025-NWTD'):
                    setv(o,'comparability','NWTD component included in Project Marinus stage costs; not an additional project total')
        elif slug == 'sydney-ring-south':
            setv(obs['SRS-AO-2026-PADR-OPT6-TIMING-STAGE1'], 'comparability', 'PADR earliest-service modelling in 2030/31 differs from the June 2026 ISP ODP, which delivers both stages in 2033-34 (AEMO 2026 Appendix A5 PDF p44, footnote 87). Keep the modelling and ODP milestone definitions separate.')
            finding('cost_aud_billion',SRS_COST,'supported_with_boundary','3.519 AUD billion nominal, combined Option 6 stages; P50 Class 5 (+100%/-50%); rounds to original 3.5')
            finding('millions_per_trunk_km',SRS_RATIO,'derived',None)
            setv(p['summary'],'analysis_conclusion',SRS_COST + ' ' + SRS_RATIO)
            setv(p['policy_chart'],'overall_status','Cost supported with rounding and price-basis qualifications. Topology, capacity and transformer evidence retained; 114 km preparatory versus approximately 100 km PADR scope remains unreconciled. ' + SRS_RATIO)
            for w in p['warnings']:
                if w['finding'] == SRS_COST and w['code'] in ('CONTRADICTION','PRICE_BASIS_UNKNOWN'): setv(w,'code','NON_COMPARABLE_SERIES'); setv(w,'implication','Retain rounding, nominal/modelled price bases and estimate uncertainty; no numerical cost contradiction is established.')
        elif slug == 'central-west-orana':
            finding('date',CWO_TIME,'supported_with_boundary','December 2028 committed-project milestone; AER staged infrastructure completion April 2028-April 2030; expansion March 2033')
            finding('feeder_lines_km',CWO_LENGTH,'supported_with_boundary','Around 150 km H2P corridor / 150 km generator-connection lines; unresolved 166 km conditional/vintage alternative')
            setv(fs['cost_aud_billion'],'supported_value',None)
            setv(fs['feeder_lines_km'],'selected_observation','CWO-AO-2025-H2P-150KM')
            setv(fs['syncons_mva'],'supported_value','Up to 1,750 MVA authorised Main system strength; AER 7 × 250 MVA units versus PACR 7 × 240 MVA units, unresolved rating discrepancy; not confirmed installed total')
            for w in p['warnings']:
                if '2029' in w['finding'] and ('timing' in w['finding'] or '2028' in w['finding']): setv(w,'finding',CWO_TIME); setv(w,'implication','Define the milestone before comparing dates; 2029 is within the AER staged window.')
        elif slug == 'project-energyconnect':
            finding('millions_per_trunk_km',PEC_RATIO,'derived',None)
            finding('lines',PEC_LINES,'supported_with_boundary',None)
            finding('feeder_lines_km',PEC_FEEDER,'supported_with_boundary','Approximately 22 km 220 kV Buronga-Red Cliffs works; feeder classification unresolved')
            setv(fs['syncons_mva'],'supported_value','2 × 100 MVA per synchronous condenser at Dinawan; additional 1 × 100 MVA at Buronga; not a certified project-wide total')
            for w in p['warnings']:
                if 'feeder' in w['finding'].lower() and ('zero' in w['finding'].lower() or '22 km' in w['finding']): setv(w,'finding',PEC_FEEDER); setv(w,'implication','Define the feeder category; preserve the source-defined lower-voltage works, without asserting a REZ-feeder total.')
        elif slug.startswith('vni-west'):
            setv(fs['cost_aud_billion'],'supported_value',None)
            if slug == 'vni-west-nsw': setv(fs['transformers'],'supported_value','3 × 1,500 MVA per Dinawan transformer in N5 scope; only 2 units included in reported VNI cost; third separately modelled; no silent total')
            else: setv(fs['transformers'],'supported_value','2 × 1,000 MVA per Tragowel transformer; excludes Dinawan units; chart 2,000 MVA is author arithmetic, not a source-stated total')
        elif slug == 'hunter-transmission-project':
            setv(fs['cost_aud_billion'],'supported_value','1,364 AUD million; CNSW-SNW Option 1 including Waratah Super Battery and related upgrades (AEMO 2026 p39 footnote79); Class 5b ±50%, dollar year unstated; not the 110 km line alone')
            setv(fs['cost_aud_billion'],'evidence','AEMO 2026 p39 reports $1,364m, Class 5b ±50%, for CNSW-SNW Option 1. Footnote79 includes Waratah Super Battery and related upgrades; do not reduce the cost boundary to the 110 km line alone. The original $1.36bn is consistent with rounding, but the dollar year is unstated. The 2024 $1,099m and 2023 approximately $1bn observations are separate non-comparable vintages.')
        elif slug == 'western-renewables-link':
            setv(obs['WRL-AO-2026-COMPLETION-NOV2029'], 'value', 'November 2029')
            setv(fs['cost_aud_billion'],'supported_value',None)
            setv(fs['cost_aud_billion'],'evidence',fs['cost_aud_billion']['evidence'].replace('supported_value is in AUD million, not the chart\'s AUD billion; no conversion is asserted.', 'No standalone replacement is established; $315m is the included WRL modification increment, not standalone WRL capex.'))
            setv(fs['network_capacity_gw'],'supported_value','More than 3,000 MW proponent carrying capability, undefined basis; not exactly 3 GW guaranteed transfer')
            setv(fs['network_capacity_gw'],'evidence',fs['network_capacity_gw']['evidence'].replace('supported_value is in MW.', 'The lower bound and capability definition must be retained.'))
        # Audit every non-null scalar: explicit unit plus the source/boundary explanation.
        units={'cost_aud_billion':'AUD billion','network_capacity_gw':'GW','trunk_km':'km','millions_per_trunk_km':'AUD million per km (chart-derived)','circuits':'circuits','lines':'lines','syncons_mva':'MVA','hubs':'hubs','transformers':'MVA (chart arithmetic, not unit count)','feeder_lines_km':'km'}
        for f in fs.values():
            value=f['supported_value']
            if isinstance(value,(int,float)):
                unit=units[f['field']]
                if slug in ('humelink','hunter-transmission-project') and f['field']=='cost_aud_billion': unit='AUD million'
                if slug in ('humelink','hunter-transmission-project','vni-west-nsw') and f['field']=='network_capacity_gw': unit='MW'
                setv(f,'supported_value',f'{value} {unit}. Boundary/qualification: {f["evidence"]}')
    for source in pack['sources']:
        source['manifest_included'] = False
    guidance = 'Source manifest paths are upstream provenance references, not files shipped in this package (manifest_included=false). Use local_path, derivative_path and sha256 for the included archives; obtain the full repository to resolve local files.'
    if guidance not in pack['agent_instructions']:
        pack['agent_instructions'].append(guidance)
    pack['generated']['path_base']='./'
    pack['generated']['audit_correction_date']='2026-09-07'
    pack['generated']['revision_note']='Archived-evidence audit corrections only; research cutoff remains 2026-09-04. Upstream build_report input hashes describe the historical upstream build, not this corrected package; referenced manifests are not asserted to be included.'
    return replacements


def text(v):
    if v is None:return ''
    if isinstance(v,(list,dict)):return json.dumps(v,ensure_ascii=False)
    return str(v)

def markdown(p):
    def table(headers, rows):
        def cell(v):return text(v).replace('|','\\|').replace('\n',' ')
        return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(cell,row))+' |' for row in rows])
    lines=['# '+p['summary']['name'],'','Research cutoff: 2026-09-04. Archived-evidence audit correction: 2026-09-07 (not new research).','','## Identity and status','']
    lines += ['**'+k.replace('_',' ').capitalize()+'.** '+text(v)+'\n' for k,v in p['summary'].items()]
    lines += ['## Policy Paper Chart comparison','', 'Original row (immutable comparison):', '', '```json',json.dumps(p['policy_chart']['original_row'],ensure_ascii=False,indent=2),'```','',p['policy_chart']['overall_status'],'']
    fields=['field','original','status','supported_value','evidence','selected_observation','source_ids','pages']
    lines += [table(fields,[[f.get(k) for k in fields] for f in p['policy_chart']['findings']]),'']
    for key in ('observations','components','timeline','warnings'):
        fields=list(dict.fromkeys(k for item in p[key] for k in item))
        lines += ['## '+key.capitalize(),'',table(fields,[[item.get(k) for k in fields] for item in p[key]]),'']
    records=json.loads((SITE/'families'/f'{p["family_slug"]}.json').read_text())['sources']
    fields=['source_id','publisher','title','publication_date','derivative_path','sha256']
    lines += ['## Sources','','Paths are relative to site/; no source manifests are asserted to be included.', '',table(fields,[[item.get(k) for k in fields] for item in records]),'']
    return '\n'.join(line.rstrip() for line in '\n'.join(lines).splitlines()).rstrip()+'\n'


def update_workbook(wb, before, pack, replacements):
    # Replace changed research strings wherever replicated, never the original chart A:N.
    # Longest strings first avoids shorter nested strings masking exact research matches.
    repl=sorted(((a,b) for a,b in replacements.items() if isinstance(b,str) and len(a)>30),key=lambda x:-len(x[0]))
    for sheet in wb:
        for row in sheet:
            for c in row:
                if sheet.title=='Policy Paper Chart' and c.column<=14:continue
                if sheet.title=='Chart Review' and c.column==4:continue
                if isinstance(c.value,str):
                    s=c.value
                    for a,b in repl:
                        if a in s:s=s.replace(a,b)
                    if s!=c.value:c.value=s
    byname={p['summary']['name']:p for p in pack['projects']}
    for row in wb['Chart Review'].iter_rows(min_row=5):
        if row[0].value not in byname:continue
        p=byname[row[0].value]; f=next(f for f in p['policy_chart']['findings'] if f['field']==row[2].value)
        for col,key in [(5,'status'),(6,'supported_value'),(7,'evidence'),(8,'source_ids'),(9,'pages'),(10,'selected_observation')]:row[col-1].value=f[key]
        row[3].value=text(f['original']);row[3].data_type='s'
    projects={p['family_slug']:p for p in pack['projects']}
    for row in wb['Projects'].iter_rows(min_row=5):
        if row[0].value not in projects:continue
        p=projects[row[0].value]
        for col,key in [(3,'name'),(4,'identity'),(5,'current_status'),(6,'delivery_status'),(7,'owner'),(8,'funding_pathway'),(9,'capacity'),(10,'schedule'),(11,'trunk_voltage_kv'),(12,'completion_date'),(13,'dependent_projects'),(14,'analysis_conclusion')]:row[col-1].value=p['summary'].get(key)
        row[14].value=p['policy_chart']['overall_status']
    selected={f['selected_observation'] for p in pack['projects'] for f in p['policy_chart']['findings'] if f['selected_observation']}
    for row in wb['Observations'].iter_rows(min_row=5):
        if row[2].value:row[17].value='selected' if row[2].value in selected else None
    mappings={'Observations':('observations',3,{2:'project_id',3:'id',4:'measure',5:'date',6:'label',7:'value',8:'unit',9:'price_basis',10:'estimate_class',11:'uncertainty',12:'status',13:'boundary',14:'comparability',15:'source_id',16:'pages',17:'excerpt',19:'notes'}), 'Components':('components',2,{2:'id',3:'stage',4:'category',5:'description',6:'boundary',7:'source_id',8:'pages'})}
    for sheet,(key,idcol,cols) in mappings.items():
        records={r['id']:r for p in pack['projects'] for r in p[key]}
        for row in wb[sheet].iter_rows(min_row=5):
            r=records.get(row[idcol-1].value)
            if r:
                for col,k in cols.items():row[col-1].value=r.get(k)
    for sheet,key,cols in [('Timeline','timeline',['date','event','status','detail','source_id','pages']),('Warnings','warnings',['severity','code','finding','implication','source_ids','pinpoints'])]:
        rows=[(p,r) for p in pack['projects'] for r in p[key]]
        for n,(p,r) in enumerate(rows,5):
            wb[sheet].cell(n,1,p['summary']['name'])
            for col,k in enumerate(cols,2):wb[sheet].cell(n,col).value=r.get(k)
    chart=wb['Policy Paper Chart']
    for row in chart.iter_rows(min_row=5):
        p=next((p for p in pack['projects'] if p['policy_chart']['original_row'].get('project')==row[0].value),None)
        if p:row[14].value=p['policy_chart']['overall_status']
    wb['Sources']['A2']='Archived source register; hashes retained from the verified upstream register. Referenced upstream manifest files are not asserted to be included. Research cutoff 2026-09-04; audit correction 2026-09-07 is not new research.'


def changes(a,b,path=''):
    if isinstance(a,dict) and isinstance(b,dict):
        for k in sorted(a.keys()|b.keys()):yield from changes(a.get(k),b.get(k),path+'/'+k)
    elif isinstance(a,list) and isinstance(b,list) and len(a)==len(b):
        for i,(x,y) in enumerate(zip(a,b)):yield from changes(x,y,path+'/'+str(i))
    elif a!=b:yield {'path':path,'before':a,'after':b}


def validate(pack,wb):
    assert pack['generated']['research_cutoff']=='2026-09-04'
    assert pack['generated']['path_base']=='./'
    count=0
    for p in pack['projects']:
        family=json.loads((SITE/'families'/f'{p["family_slug"]}.json').read_text())
        for k in ('summary','policy_chart','observations','components','timeline','warnings'):assert family[k]==p[k],(p['family_slug'],k)
        assert (SITE/'families'/f'{p["family_slug"]}.md').read_text()==markdown(p)
        for f in p['policy_chart']['findings']:
            count+=1
            assert f['supported_value'] is None or isinstance(f['supported_value'],str)
            rows=[r for r in wb['Chart Review'].iter_rows(min_row=5) if r[0].value==p['summary']['name'] and r[2].value==f['field']]
            assert len(rows)==1
            for c,k in [(5,'status'),(6,'supported_value'),(7,'evidence'),(8,'source_ids'),(9,'pages'),(10,'selected_observation')]:assert (rows[0][c-1].value or None)==(f[k] or None),(p['family_slug'],f['field'],k)
        narrative=json.dumps({k:v for k,v in p.items() if k!='policy_chart'})
        if p['family_slug']=='sydney-ring-south':assert not re.search(r'3[,.]510|transpos|correct the transcription|unverified transcription',narrative)
        project_row=next(r for r in wb['Projects'].iter_rows(min_row=5) if r[0].value==p['family_slug'])
        assert project_row[13].value==p['summary']['analysis_conclusion']
        assert project_row[14].value==p['policy_chart']['overall_status']
        chart_row=next(r for r in wb['Policy Paper Chart'].iter_rows(min_row=5) if r[0].value==p['policy_chart']['original_row']['project'])
        assert chart_row[14].value==p['policy_chart']['overall_status']
        for key,sheet,idcol,cols in [('observations','Observations',3,{6:'label',7:'value',8:'unit',9:'price_basis',13:'boundary',14:'comparability',17:'excerpt',19:'notes'}),('components','Components',2,{3:'stage',4:'category',5:'description',6:'boundary'})]:
            rows={r[idcol-1].value:r for r in wb[sheet].iter_rows(min_row=5)}
            for item in p[key]:
                for c,k in cols.items():assert (rows[item['id']][c-1].value or None)==(item.get(k) or None),(sheet,item['id'],k)
    for key,sheet,cols in [('timeline','Timeline',['date','event','status','detail','source_id','pages']),('warnings','Warnings',['severity','code','finding','implication','source_ids','pinpoints'])]:
        for n,item in enumerate([r for p in pack['projects'] for r in p[key]],5):
            for c,k in enumerate(cols,2):assert (wb[sheet].cell(n,c).value or None)==(item.get(k) or None),(sheet,n,k)
    candidate=copy.deepcopy(pack);correct(candidate);assert candidate==pack,'Correction must be idempotent'
    selected={f['selected_observation'] for p in pack['projects'] for f in p['policy_chart']['findings'] if f['selected_observation']}
    for row in wb['Observations'].iter_rows(min_row=5):
        if row[2].value:assert row[17].value==('selected' if row[2].value in selected else None)
    assert count==120
    assert wb['Chart Review']['D45'].data_type=='s' and wb['Chart Review']['D45'].value=='=C11/E11 *1000'
    print('PASS: 120 unit/boundary-qualified findings; ten coherent family JSON/Markdown exports; workbook Chart Review; literal D45; unchanged cutoff')


def main():
    args=argparse.ArgumentParser();args.add_argument('--check',action='store_true');opts=args.parse_args()
    path=SITE/'evidence-pack.json';pack=json.loads(path.read_text());wb=openpyxl.load_workbook(SITE/'evidence-workbook.xlsx')
    if opts.check:
        validate(pack,wb)
        log=json.loads(LOG.read_text());assert protected(pack,wb)==log['protected_before']==log['protected_after']
        print('PASS: before/after immutable original rows, original chart inputs/styles/formulas, excerpts and all archived source bytes')
        return
    before=copy.deepcopy(pack);guard=protected(pack,wb);styles=[(s.title,c.coordinate,str(c._style)) for s in wb for row in s for c in row]
    replacements=correct(pack)
    update_workbook(wb,before,pack,replacements)
    assert protected(pack,wb)==guard
    assert styles==[(s.title,c.coordinate,str(c._style)) for s in wb for row in s for c in row]
    path.write_text(json.dumps(pack,ensure_ascii=False,indent=2)+'\n')
    for p in pack['projects']:
        fp=SITE/'families'/f'{p["family_slug"]}.json';family=json.loads(fp.read_text())
        for k in ('summary','policy_chart','observations','components','timeline','warnings'):family[k]=p[k]
        for k in ('family_slug','project_id','source_ids'):family.pop(k,None)
        source_by_id = {s['source_id']: s for s in pack['sources']}
        family['sources'] = [source_by_id[s['source_id']] for s in family['sources']]
        family['build']['path_base']='../';family['build']['audit_correction_date']='2026-09-07';family['build']['revision_note']=pack['generated']['revision_note']
        fp.write_text(json.dumps(family,ensure_ascii=False,indent=2)+'\n')
        fp.with_suffix('.md').write_text(markdown(p))
    wb.save(SITE/'evidence-workbook.xlsx')
    delta=list(changes(before,pack))
    if delta:
        prior=json.loads(LOG.read_text()) if LOG.exists() else None
        if prior:
            assert prior['protected_before']==guard
            delta=prior['changes']+delta
        log={'correction_date':'2026-09-07','research_cutoff':'2026-09-04','protected_before':guard,'protected_after':protected(pack,wb),'changes':delta,'supported_value_audit':[{'family':p['family_slug'],'field':f['field'],'observation':f['selected_observation'],'source_ids':f['source_ids'],'pages':f['pages'],'supported_value':f['supported_value']} for p in pack['projects'] for f in p['policy_chart']['findings']], 'source_pinpoints':{'R1':'ML-TN-2025-RITT PDF pp6,18, text205-218,806-825; HUM-AEMO-2026-ISP-A5 p59 text2299-2335','R2':'HUM-AEMO-2026-ISP-A5 pp17,52-53 text591-593,2003,2030-2031,2054; VNV-TRANSGRID-2023-PACR p48; individual finding references in supported_value_audit','R3':'SRS-TG-2026-PADR PDF p38 text3068-3072; 2023 preparatory p8 versus PADR p43','R4':'VNV-TRANSGRID-2023-PACR p48 text2072-2079; HUM-AEMO-2026-ISP-A5 p52','R5':'CWO-18 p11 text475-477; CWO-26 line18; CWO-23 p10 text279-280 and conditional scope351-352,625-626','R6':'PEC-TG-2026-COMPLETION text161-166; PEC-TG-PROJECT-PAGE text1705; AEMO2024 pp17-18 for line/feeder definitions'},'changed_export_paths':['site/evidence-pack.json','site/evidence-workbook.xlsx']+[f'site/families/{p["family_slug"]}.{ext}' for p in pack['projects'] for ext in ['json','md']], 'workbook_note':'All cell styles preserved; Policy Paper Chart A:N values/formulas unchanged; research strings refreshed across all sheets; D45 literal audit text. No source or source-manifest generation.'}
        LOG.write_text(json.dumps(log,ensure_ascii=False,indent=2)+'\n')
    if LOG.exists():
        log=json.loads(LOG.read_text())
        log['source_pinpoints']['R7'] = 'SRS-TG-2026-PADR PDF p38 earliest 2030/31; HUM-AEMO-2026-ISP-A5 PDF p44 footnote87 (text1656-1658) ODP2033-34'
        log['source_pinpoints']['R8'] = 'HUM-AEMO-2026-ISP-A5 PDF p24 explicitly states expected completion November2029; existing WRL excerpt and notes preserved'
        log['record_ids_by_json_path']={f'/projects/{i}/{key}/{j}':item.get('id',item.get('field',item.get('event',item.get('code')))) for i,p in enumerate(pack['projects']) for key in ('observations','components','timeline','warnings') for j,item in enumerate(p[key])}
        log['family_by_json_path']={f'/projects/{i}':p['family_slug'] for i,p in enumerate(pack['projects'])}
        LOG.write_text(json.dumps(log,ensure_ascii=False,indent=2)+'\n')
    validate(pack,openpyxl.load_workbook(SITE/'evidence-workbook.xlsx'))
    print('PASS: source bytes, original chart values/formulas/styles and raw excerpts unchanged; correction log written')

if __name__=='__main__':main()
